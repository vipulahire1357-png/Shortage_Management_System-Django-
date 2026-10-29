"""
Export Material Shortage Register workbooks for individual projects.

Generates one Excel workbook per project saved under:
    Material Shortage Registers/<project_id>/<project_id>_Net_shortage_Table.xlsx

Features:
- Preserves numeric values as numbers and Material Codes as text.
- Places Vendor Detail after Free Stock and Final Date before Remarks.
- Creates an Excel Table with filters and frozen header row.
- Configurable output directory.
"""

from __future__ import annotations

import io
import os
from pathlib import Path
import re
import sqlite3
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = Path(os.getenv("ABB_DATA_ROOT", PROJECT_ROOT / "data")).resolve()
SERVICES_DB_PATH = Path(os.getenv("ABB_SERVICES_DB", DATA_DIR / "services.db")).resolve()
DEFAULT_SHORTAGE_REGISTERS_DIR = Path(
    os.getenv("SHORTAGE_REGISTERS_DIR", str(PROJECT_ROOT / "Material Shortage Registers"))
)

# Ordered headers for the Material Shortage Register Excel file
SHORTAGE_REGISTER_HEADERS = [
    "Material Ref Code",
    "WBS Element",
    "Material Code",
    "Material Description",
    "ZPP7 Quantity",
    "Project Stock Quantity",
    "PO Quantity",
    "Open PO Quantity",
    "PR Quantity",
    "Net Shortage",
    "Free Stock",
    "Vendor Detail",
    "Final Date",
    "Remarks",
    "Scope",
]

NUMERIC_COLUMNS = {
    "ZPP7 Quantity",
    "Project Stock Quantity",
    "PO Quantity",
    "Open PO Quantity",
    "PR Quantity",
    "Net Shortage",
    "Free Stock",
}

TEXT_FORMAT_COLUMNS = {
    "Material Ref Code",
    "WBS Element",
    "Material Code",
}

COLUMN_WIDTHS = [
    22,  # Material Ref Code
    22,  # WBS Element
    22,  # Material Code
    40,  # Material Description
    14,  # ZPP7 Quantity
    18,  # Project Stock Quantity
    14,  # PO Quantity
    16,  # Open PO Quantity
    14,  # PR Quantity
    14,  # Net Shortage
    14,  # Free Stock
    30,  # Vendor Detail
    16,  # Final Date
    35,  # Remarks
    14,  # Scope
]


def build_material_shortage_workbook(
    rows: list[dict[str, Any] | sqlite3.Row],
    table_name_suffix: str = "Register",
) -> openpyxl.Workbook:
    """Build an openpyxl Workbook containing formatted Material Shortage Register records."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Material Shortage Register"

    # Header styling (ABB styling)
    header_fill = PatternFill("solid", fgColor="B5121B")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append(SHORTAGE_REGISTER_HEADERS)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    for row in rows:
        row_data = []
        for col_name in SHORTAGE_REGISTER_HEADERS:
            val = row[col_name] if isinstance(row, sqlite3.Row) else row.get(col_name)
            if col_name in NUMERIC_COLUMNS:
                try:
                    val = float(val) if val is not None else 0.0
                except (TypeError, ValueError):
                    val = 0.0
            else:
                val = str(val) if val is not None else ""
            row_data.append(val)
        ws.append(row_data)

    total_rows = len(rows)
    total_cols = len(SHORTAGE_REGISTER_HEADERS)
    max_col_letter = get_column_letter(total_cols)

    # Format text cells (especially Material Code to preserve leading zeros)
    for r_idx in range(2, total_rows + 2):
        for c_idx, col_name in enumerate(SHORTAGE_REGISTER_HEADERS, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            if col_name in TEXT_FORMAT_COLUMNS:
                cell.number_format = "@"

    # Freeze header row
    ws.freeze_panes = "A2"

    # Set column widths
    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Add real Excel Table with filters
    safe_suffix = re.sub(r"[^A-Za-z0-9_]", "_", table_name_suffix)
    if not safe_suffix or not safe_suffix[0].isalpha():
        safe_suffix = "T_" + safe_suffix
    table_name = f"Shortage_{safe_suffix}"[:30]

    # Excel tables require at least header + 1 data row
    table_ref = f"A1:{max_col_letter}{max(2, total_rows + 1)}"
    tab = Table(displayName=table_name, ref=table_ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tab)

    return wb


def export_material_shortage_register(
    project_id: str,
    db_path: Path = SERVICES_DB_PATH,
    output_dir: Path = DEFAULT_SHORTAGE_REGISTERS_DIR,
) -> dict[str, Any]:
    """
    Export the current project's Material Shortage Register to Excel.
    Saves to: <output_dir>/<project_id>/<project_id>_Net_shortage_Table.xlsx
    """
    clean_id = project_id.strip()
    if not clean_id or not re.match(r"^[A-Za-z0-9][A-Za-z0-9\-\.]{0,63}$", clean_id):
        raise ValueError(f"Invalid project ID {project_id!r}")

    if not db_path.is_file():
        raise FileNotFoundError(f"Services database not found: {db_path}")

    # Nested folder: Material Shortage Registers/<project_id>/
    project_dir = output_dir / clean_id
    project_dir.mkdir(parents=True, exist_ok=True)
    out_file = project_dir / f"{clean_id}_Net_shortage_Table.xlsx"

    with sqlite3.connect(f"file:{db_path.as_posix()}?mode=ro", uri=True) as conn:
        conn.row_factory = sqlite3.Row
        # Verify net_shortage exists
        tables = {r[0].casefold(): r[0] for r in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table'"
        )}
        table_name = tables.get("net_shortage")
        if not table_name:
            raise RuntimeError("net_shortage table not found. Run shortage processing first.")

        # Read only this project's rows, ordered for the register
        query = (
            f'SELECT * FROM "{table_name}" '
            f'WHERE "project_id" = ? '
            f'ORDER BY "Material Code", "WBS Element"'
        )
        rows = conn.execute(query, (clean_id,)).fetchall()

    # Never fall back to another project's rows. An empty project export stays
    # empty instead of silently including unrelated project data.

    wb = build_material_shortage_workbook(rows, table_name_suffix=clean_id)
    wb.save(out_file)

    return {
        "project_id": clean_id,
        "filename": out_file.name,
        "path": str(out_file.resolve()),
        "rows": len(rows),
    }
