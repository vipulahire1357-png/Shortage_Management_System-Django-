"""Read and update vendor date registers used by the Vendor Response page."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import sqlite3

import openpyxl


SUPPORTED_WORKBOOKS = {".xlsx", ".xlsm"}
REQUIRED_COLUMNS = (
    "Project ID", "Vendor Detail", "WBS Element", "Material Code",
    "Description", "Quantity", "Vendor Date", "Proposed Date", "Final Date",
)


def _workbook_files(folder: Path) -> list[Path]:
    return [
        path for path in folder.rglob("*")
        if path.is_file()
        and path.suffix.casefold() in SUPPORTED_WORKBOOKS
        and not path.name.startswith("~$")
    ]


def list_vendor_folders(registers_dir: Path) -> list[str]:
    """List vendor folders and legacy root-level workbooks as selectable entries."""
    root = Path(registers_dir).resolve()
    if not root.is_dir():
        return []

    names = set()
    for entry in root.iterdir():
        if entry.is_dir() and _workbook_files(entry):
            names.add(entry.name)
        elif entry.is_file() and entry.suffix.casefold() in SUPPORTED_WORKBOOKS and not entry.name.startswith("~$"):
            names.add(entry.stem)
    return sorted(names, key=str.casefold)


def resolve_vendor_workbook(registers_dir: Path, vendor_folder: str) -> Path:
    """Resolve a vendor folder to its newest workbook without allowing path escape."""
    root = Path(registers_dir).resolve()
    if not vendor_folder or Path(vendor_folder).name != vendor_folder or vendor_folder in {".", ".."}:
        raise ValueError("Select a valid vendor folder.")

    folder = (root / vendor_folder).resolve()
    try:
        folder.relative_to(root)
    except ValueError as exc:
        raise ValueError("The selected vendor folder is outside Vendor Date Registers.") from exc

    if folder.is_dir():
        candidates = _workbook_files(folder)
    else:
        candidates = [
            path for path in root.iterdir()
            if path.is_file()
            and path.suffix.casefold() in SUPPORTED_WORKBOOKS
            and path.stem.casefold() == vendor_folder.casefold()
        ] if root.is_dir() else []

    safe_candidates = []
    for path in candidates:
        resolved = path.resolve()
        try:
            resolved.relative_to(root)
        except ValueError:
            continue
        safe_candidates.append(resolved)
    if not safe_candidates:
        raise FileNotFoundError(f"No Excel workbook was found for vendor folder {vendor_folder}.")
    return max(safe_candidates, key=lambda path: path.stat().st_mtime_ns)


def _header_columns(worksheet) -> dict[str, int]:
    columns = {
        str(cell.value).strip().casefold(): cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }
    missing = [name for name in REQUIRED_COLUMNS if name.casefold() not in columns]
    if missing:
        raise ValueError(f"Vendor workbook is missing required columns: {', '.join(missing)}")
    return columns


def _read_date(value) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for date_format in (
        "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y",
        "%m/%d/%Y", "%d/%m/%y", "%d-%m-%y", "%d-%b-%Y", "%d-%b-%y",
    ):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        return None


def _date_text(value) -> str:
    parsed = _read_date(value)
    return parsed.isoformat() if parsed else (str(value).strip() if value is not None else "")


def read_vendor_register(workbook_path: Path) -> list[dict]:
    workbook_path = Path(workbook_path)
    workbook = openpyxl.load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        worksheet = workbook.active
        columns = _header_columns(worksheet)
        result = []
        for row_number, cells in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            values = {name: cells[columns[name.casefold()] - 1] for name in REQUIRED_COLUMNS}
            if not any(value is not None and str(value).strip() for value in values.values()):
                continue
            vendor_date = _date_text(values["Vendor Date"])
            proposed_date = _date_text(values["Proposed Date"])
            result.append({
                "row_number": row_number,
                "project_id": str(values["Project ID"] or "").strip(),
                "vendor": str(values["Vendor Detail"] or "").strip(),
                "wbs_element": str(values["WBS Element"] or "").strip(),
                "material_code": str(values["Material Code"] or "").strip(),
                "description": str(values["Description"] or "").strip(),
                "quantity": values["Quantity"],
                "vendor_date": vendor_date,
                "vendor_date_valid": _read_date(values["Vendor Date"]) is not None,
                "proposed_date": proposed_date,
                "final_date": _date_text(values["Final Date"]),
            })
        return result
    finally:
        workbook.close()


def save_proposed_date(workbook_path: Path, row_number: int, proposed_date: str) -> None:
    workbook_path = Path(workbook_path)
    workbook = openpyxl.load_workbook(
        workbook_path,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        worksheet = workbook.active
        columns = _header_columns(worksheet)
        if row_number < 2 or row_number > worksheet.max_row:
            raise ValueError("The selected vendor register row no longer exists.")
        value = proposed_date.strip()
        parsed = _read_date(value) if value else None
        if value and parsed is None:
            raise ValueError("Enter Proposed Date as YYYY-MM-DD or another common date format.")
        cell = worksheet.cell(row=row_number, column=columns["proposed date"])
        cell.value = parsed
        if parsed:
            cell.number_format = "yyyy-mm-dd"
        workbook.save(workbook_path)
    finally:
        workbook.close()


def apply_accepted_vendor_dates(
    workbook_path: Path,
    db_path: Path,
    registers_dir: Path,
) -> dict:
    """Apply every valid Vendor Date to matching rows and regenerate project registers."""
    workbook_path = Path(workbook_path)
    workbook = openpyxl.load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        worksheet = workbook.active
        columns = _header_columns(worksheet)
        chosen = []
        for row_number, cells in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            values = {name: cells[columns[name.casefold()] - 1] for name in REQUIRED_COLUMNS}
            project_id = str(values["Project ID"] or "").strip()
            vendor = str(values["Vendor Detail"] or "").strip()
            wbs = str(values["WBS Element"] or "").strip()
            material = str(values["Material Code"] or "").strip()
            vendor_date = _read_date(values["Vendor Date"])
            if project_id and vendor and wbs and material and vendor_date:
                chosen.append((row_number, project_id, vendor, wbs, material, vendor_date.isoformat()))
    finally:
        workbook.close()

    if not chosen:
        raise ValueError("No rows with valid Vendor Dates and matching project, vendor, WBS, and material values were found.")
    if not Path(db_path).is_file():
        raise FileNotFoundError(f"Services database was not found: {db_path}")

    affected_projects = set()
    missing_rows = []
    updated_register_rows = 0
    conn = sqlite3.connect(db_path)
    try:
        tables = {row[0].casefold() for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if "net_shortage" not in tables:
            raise RuntimeError("The net_shortage table is not available.")
        has_vendor_track = "vendor_track" in tables
        for row_number, project_id, vendor, wbs, material, final_date in chosen:
            cursor = conn.execute(
                'UPDATE net_shortage SET "Final Date"=? WHERE "project_id"=? AND "Vendor Detail"=? AND "WBS Element"=? AND "Material Code"=?',
                (final_date, project_id, vendor, wbs, material),
            )
            if cursor.rowcount == 0:
                missing_rows.append(row_number)
                continue
            updated_register_rows += cursor.rowcount
            affected_projects.add(project_id)
            if has_vendor_track:
                conn.execute(
                    'UPDATE vendor_track SET "Final Date"=? WHERE "project_id"=? AND "Vendor Detail"=? AND "WBS Element"=? AND "Material Code"=?',
                    (final_date, project_id, vendor, wbs, material),
                )
        conn.commit()
    finally:
        conn.close()

    regenerated = []
    export_errors = []
    if affected_projects:
        from services.shortage_export import export_material_shortage_register
        for project_id in sorted(affected_projects):
            try:
                export_material_shortage_register(project_id, db_path=Path(db_path), output_dir=Path(registers_dir))
                regenerated.append(project_id)
            except Exception as exc:
                export_errors.append(f"{project_id}: {exc}")

    return {
        "selected_rows": len(chosen),
        "updated_register_rows": updated_register_rows,
        "projects": regenerated,
        "missing_rows": missing_rows,
        "export_errors": export_errors,
    }
