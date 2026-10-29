"""
Vendor Date Request Excel Export Service.

Features:
- Exports filtered Vendor Track records to an Excel workbook.
- Evaluates First/Second remark rule across all rows in the export batch:
    * First: All three date cells (Vendor Date, Proposed Date, Final Date) are empty across all rows.
    * Second: At least one date cell across any row is populated.
- Saves under:
    Vendor Registers/<Vendor_Folder>/Date Request/<Remark>.xlsx
- Versioning scheme:
    Preserves previous exports for traceability (First.xlsx -> First_v2.xlsx, etc.).
- Requires single-vendor selection; rejects empty datasets and multi-vendor batches.
- Excel usability: Table with filters, frozen header row, numeric quantities, text material codes.
"""

from __future__ import annotations

import os
from pathlib import Path
import re
import sqlite3
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = Path(os.getenv("ABB_DATA_ROOT", PROJECT_ROOT / "data")).resolve()
SERVICES_DB_PATH = Path(os.getenv("ABB_SERVICES_DB", DATA_DIR / "services.db")).resolve()
DEFAULT_VENDOR_REGISTERS_DIR = Path(
    os.getenv("VENDOR_REGISTERS_DIR", str(PROJECT_ROOT / "Vendor Registers"))
)

# Required Excel columns in exact order
VENDOR_REQUEST_COLUMNS = [
    "Vendor Detail",
    "WBS Element",
    "Material Code",
    "Description",
    "Quantity",
    "Vendor Date",
    "Proposed Date",
    "Final Date",
]
INTERNAL_PROJECT_ID_COLUMN = "Project ID"

NUMERIC_COLUMNS = {"Quantity"}
TEXT_COLUMNS = {"Material Code", "WBS Element", "Vendor Detail", "Description"}
DATE_COLUMNS = {"Vendor Date", "Proposed Date", "Final Date"}

COLUMN_WIDTHS = [
    32,  # Vendor Detail
    22,  # WBS Element
    22,  # Material Code
    42,  # Description
    14,  # Quantity
    16,  # Vendor Date
    16,  # Proposed Date
    16,  # Final Date
]


def sanitize_vendor_folder_name(vendor: str) -> str:
    """
    Sanitize vendor name to be safe as a directory name on Windows and other OSes.
    Removes invalid characters: \\ / : * ? " < > |
    Strips trailing spaces and periods.
    """
    if not vendor or not vendor.strip():
        return "UNKNOWN_VENDOR"
    # Replace invalid filesystem characters with a space
    clean = re.sub(r'[\\/*?:"<>|]', " ", vendor.strip())
    # Collapse multiple spaces
    clean = re.sub(r"\s+", " ", clean).strip()
    # Windows does not allow folder names ending in dot or space
    clean = clean.rstrip(". ")
    return clean or "UNKNOWN_VENDOR"


def is_date_populated(val: Any) -> bool:
    """
    Check if a date field is populated with a valid, meaningful value.
    Treats None, empty strings, whitespace-only, and null-like strings as empty.
    """
    if val is None:
        return False
    s = str(val).strip()
    if not s:
        return False
    if s.lower() in ("none", "null", "nan", "nat", "—", "-"):
        return False
    return True


def calculate_batch_remark(rows: list[dict[str, Any] | sqlite3.Row]) -> str:
    """
    Calculate First/Second remark for the entire batch.
    - First: Every Vendor Date, Proposed Date, and Final Date cell across all rows is empty.
    - Second: At least one of these date fields is populated in at least one row.
    """
    for row in rows:
        for date_col in DATE_COLUMNS:
            val = row[date_col] if isinstance(row, sqlite3.Row) else row.get(date_col)
            if is_date_populated(val):
                return "Second"
    return "First"


def get_versioned_export_path(directory: Path, base_name: str) -> Path:
    """
    Return a non-conflicting versioned file path:
    If <base_name>.xlsx does not exist -> <base_name>.xlsx
    If it exists -> <base_name>_v2.xlsx, <base_name>_v3.xlsx, etc.
    """
    candidate = directory / f"{base_name}.xlsx"
    if not candidate.exists():
        return candidate

    v = 2
    while True:
        candidate = directory / f"{base_name}_v{v}.xlsx"
        if not candidate.exists():
            return candidate
        v += 1


def build_vendor_request_workbook(
    rows: list[dict[str, Any] | sqlite3.Row],
    vendor_name: str,
    table_name_suffix: str = "DateRequest",
) -> openpyxl.Workbook:
    """
    Build an Excel workbook for the Vendor Date Request.
    Formatted as an Excel table with filters, frozen headers, and proper types.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vendor Date Request"

    # Header row styling (ABB theme)
    header_fill = PatternFill("solid", fgColor="B5121B")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append(VENDOR_REQUEST_COLUMNS)
    project_id_column = len(VENDOR_REQUEST_COLUMNS) + 1
    ws.cell(row=1, column=project_id_column, value=INTERNAL_PROJECT_ID_COLUMN)
    ws.column_dimensions[get_column_letter(project_id_column)].hidden = True
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    for row in rows:
        row_data = []
        for col_name in VENDOR_REQUEST_COLUMNS:
            if isinstance(row, sqlite3.Row):
                val = row[col_name]
            else:
                val = row.get(col_name)
            if col_name in NUMERIC_COLUMNS:
                try:
                    val = float(val) if val is not None else 0.0
                except (TypeError, ValueError):
                    val = 0.0
            elif col_name in DATE_COLUMNS:
                # Keep date as string or empty
                val = str(val).strip() if is_date_populated(val) else ""
            else:
                val = str(val) if val is not None else ""
            row_data.append(val)
        ws.append(row_data)
        project_id = row["project_id"] if isinstance(row, sqlite3.Row) else row.get("Project ID", row.get("project_id"))
        project_cell = ws.cell(row=ws.max_row, column=project_id_column, value=str(project_id or ""))
        project_cell.number_format = "@"
        project_cell.protection = Protection(locked=True)

    total_rows = len(rows)
    total_cols = len(VENDOR_REQUEST_COLUMNS)
    max_col_letter = get_column_letter(total_cols)

    # Format text cells (especially Material Code to preserve leading zeros)
    for r_idx in range(2, total_rows + 2):
        for c_idx, col_name in enumerate(VENDOR_REQUEST_COLUMNS, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            if col_name in TEXT_COLUMNS:
                cell.number_format = "@"
            cell.protection = Protection(locked=col_name not in {"Vendor Date", "Proposed Date"})

    # Freeze header row
    ws.freeze_panes = "A2"

    # Set column widths
    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Add real Excel Table with filters
    safe_suffix = re.sub(r"[^A-Za-z0-9_]", "_", table_name_suffix)
    if not safe_suffix or not safe_suffix[0].isalpha():
        safe_suffix = "V_" + safe_suffix
    table_name = f"VendorReq_{safe_suffix}"[:30]

    # Excel tables require at least header + 1 data row
    table_ref = f"A1:{max_col_letter}{max(2, total_rows + 1)}"
    tab = Table(displayName=table_name, ref=table_ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tab)

    # Protect every request field except Vendor Date and Proposed Date.
    # Keep filtering and sorting available for the Excel table.
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False

    return wb


def enforce_vendor_date_fields_editable(workbook_path: Path) -> None:
    """Update an existing request workbook so only Vendor Date and Proposed Date are editable."""
    workbook = openpyxl.load_workbook(workbook_path)
    try:
        worksheet = workbook.active
        headers = {
            worksheet.cell(row=1, column=column).value: column
            for column in range(1, worksheet.max_column + 1)
        }
        editable_columns = {headers.get("Vendor Date"), headers.get("Proposed Date")}
        editable_columns.discard(None)
        if "Vendor Date" not in headers or "Proposed Date" not in headers:
            raise ValueError("The vendor date request workbook must have Vendor Date and Proposed Date columns.")

        protection_needed = not worksheet.protection.sheet
        for column in range(1, worksheet.max_column + 1):
            should_lock = column not in editable_columns
            for row in range(2, worksheet.max_row + 1):
                if worksheet.cell(row=row, column=column).protection.locked != should_lock:
                    protection_needed = True
                    break
            if protection_needed:
                break

        if protection_needed:
            for column in range(1, worksheet.max_column + 1):
                protection = Protection(locked=column not in editable_columns)
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row, column=column).protection = protection
            worksheet.protection.sheet = True
            worksheet.protection.autoFilter = False
            worksheet.protection.sort = False
            workbook.save(workbook_path)
    finally:
        workbook.close()


def export_vendor_date_request(
    vendor: str,
    project_ids: list[str] | None = None,
    wbs_elements: list[str] | None = None,
    db_path: Path = SERVICES_DB_PATH,
    output_root: Path = DEFAULT_VENDOR_REGISTERS_DIR,
) -> dict[str, Any]:
    """
    Export the currently filtered Vendor Track records for a single vendor.
    Follows First/Second remark rule and saves to:
    Vendor Registers/<vendor>/Date Request/<Remark>.xlsx
    """
    clean_vendor = (vendor or "").strip()
    if not clean_vendor or clean_vendor in ("All Vendors", "All"):
        raise ValueError(
            "Please select a specific vendor before exporting a Vendor Date Request. "
            "Multi-vendor export is not allowed."
        )

    if not db_path.is_file():
        raise FileNotFoundError(f"Services database not found: {db_path}")

    # Build query: export only rows with a strictly positive quantity.
    conditions: list[str] = [
        '"Vendor Detail" = ?',
        '"Quantity" > 0',
    ]
    params: list[Any] = [clean_vendor]

    if project_ids:
        valid_pids = [p.strip() for p in project_ids if p.strip()]
        if valid_pids:
            placeholders = ",".join("?" for _ in valid_pids)
            conditions.append(f'"project_id" IN ({placeholders})')
            params.extend(valid_pids)

    if wbs_elements:
        valid_wbs = [w.strip() for w in wbs_elements if w.strip()]
        if valid_wbs:
            placeholders = ",".join("?" for _ in valid_wbs)
            conditions.append(f'"WBS Element" IN ({placeholders})')
            params.extend(valid_wbs)

    where = " WHERE " + " AND ".join(conditions)

    with sqlite3.connect(f"file:{db_path.as_posix()}?mode=ro", uri=True) as conn:
        conn.row_factory = sqlite3.Row
        tables = {r[0].casefold(): r[0] for r in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table'"
        )}
        if "vendor_track" not in tables:
            raise RuntimeError("vendor_track table not found. Process project registers first.")

        query = f'''
            SELECT
                "Vendor Detail",
                "WBS Element",
                "Material Code",
                "Description",
                "Quantity",
                "Vendor Date",
                "Proposed Date",
                "Final Date",
                "project_id"
            FROM "vendor_track"{where}
            ORDER BY "WBS Element", "Material Code"
        '''
        rows = conn.execute(query, params).fetchall()

    if not rows:
        raise ValueError(
            f"No records found for vendor '{clean_vendor}' with the current filters. "
            "Cannot export an empty request."
        )

    # Double check all rows belong to the requested vendor
    distinct_vendors = {r["Vendor Detail"] for r in rows if r["Vendor Detail"]}
    if len(distinct_vendors) > 1:
        raise ValueError(
            f"Multiple vendors found in filtered records ({len(distinct_vendors)}). "
            "Select a single vendor before exporting."
        )

    # Calculate First/Second remark
    remark = calculate_batch_remark(rows)

    # Destination folder: Vendor Registers/<Vendor>/Date Request/
    folder_name = sanitize_vendor_folder_name(clean_vendor)
    dest_dir = Path(output_root) / folder_name / "Date Request"
    dest_dir.mkdir(parents=True, exist_ok=True)

    # Versioned file path
    target_file = get_versioned_export_path(dest_dir, remark)

    # Build and save Excel workbook
    wb = build_vendor_request_workbook(rows, vendor_name=clean_vendor, table_name_suffix=folder_name[:15])
    wb.save(target_file)

    # Distinct projects present in the exported records
    actual_projects = sorted({r["project_id"] for r in rows if r["project_id"]})
    wbs_display = (
        sorted(list(set(wbs_elements)))
        if wbs_elements
        else "All WBS Elements"
    )

    try:
        relative_path = str(target_file.relative_to(PROJECT_ROOT))
    except ValueError:
        relative_path = str(target_file)

    return {
        "success": True,
        "vendor": clean_vendor,
        "vendor_folder": folder_name,
        "project_count": len(actual_projects),
        "projects": actual_projects,
        "wbs_elements": wbs_display,
        "row_count": len(rows),
        "remark": remark,
        "filename": target_file.name,
        "output_path": str(target_file.resolve()),
        "relative_path": relative_path,
        "message": (
            f"Vendor Date Request exported successfully as {target_file.name} "
            f"({remark} request, {len(rows)} rows)."
        ),
    }
