'''Build vendor-wise tracking data from Material Shortage Registers and services.db.'''
from __future__ import annotations

import os
from pathlib import Path
import re
import sqlite3
import sys
from typing import Any

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = Path(os.getenv("ABB_DATA_ROOT", PROJECT_ROOT / "data")).resolve()
SERVICES_DB_PATH = Path(os.getenv("ABB_SERVICES_DB", DATA_DIR / "services.db")).resolve()
DEFAULT_SHORTAGE_REGISTERS_DIR = Path(
    os.getenv("SHORTAGE_REGISTERS_DIR", str(PROJECT_ROOT / "Material Shortage Registers"))
)
SOURCE_TABLE = "net_shortage"
OUTPUT_TABLE = "vendor_track"
OUTPUT_COLUMNS = [
    "Vendor Detail", "WBS Element", "Material Code", "Description",
    "Quantity", "Vendor Date", "Proposed Date", "Final Date", "Status",
]

# Required register fields (mapped flexibly to handle minor naming variations)
REQUIRED_REGISTER_FIELDS = {
    "vendor_detail": ["vendor detail", "vendor", "vendor name"],
    "wbs_element": ["wbs element", "wbs"],
    "material_code": ["material code", "material"],
    "description": ["material description", "description"],
    "quantity": ["net shortage", "quantity", "shortage"],
}


def quote_identifier(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def resolve_table(conn: sqlite3.Connection, name: str) -> str:
    tables = {str(r[0]).casefold(): str(r[0]) for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
    )}
    actual = tables.get(name.casefold())
    if actual is None:
        raise RuntimeError(f"Required table {name!r} not found in {SERVICES_DB_PATH}.")
    return actual


def resolve_columns(conn: sqlite3.Connection, table: str, required: list[str]) -> dict[str, str]:
    available = {str(r[1]).strip().casefold(): str(r[1]) for r in conn.execute(
        f"PRAGMA table_info({quote_identifier(table)})"
    )}
    missing = [c for c in required if c.casefold() not in available]
    if missing:
        raise RuntimeError(f"Table {table!r} is missing columns {missing}. Found: {list(available.values())}")
    return {c: available[c.casefold()] for c in required}


def clean_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _ensure_vendor_track_schema(conn: sqlite3.Connection) -> None:
    """Create vendor_track table with project_id column if it does not exist.
    If it exists but lacks project_id, add the column non-destructively."""
    tables = {str(r[0]).casefold(): str(r[0]) for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table'"
    )}
    if OUTPUT_TABLE.casefold() not in tables:
        conn.execute(
            f'CREATE TABLE {quote_identifier(OUTPUT_TABLE)} ('
            '"project_id" TEXT, '
            '"Vendor Detail" TEXT, "WBS Element" TEXT, "Material Code" TEXT, '
            '"Description" TEXT, "Quantity" REAL, "Vendor Date" TEXT, '
            '"Proposed Date" TEXT, "Final Date" TEXT, "Status" TEXT'
            ')'
        )
        return
    # Table exists - check for project_id column
    cols = {str(r[1]).casefold() for r in conn.execute(
        f'PRAGMA table_info({quote_identifier(OUTPUT_TABLE)})'
    )}
    if 'project_id' not in cols:
        conn.execute(
            f'ALTER TABLE {quote_identifier(OUTPUT_TABLE)} ADD COLUMN "project_id" TEXT'
        )


def validate_material_shortage_register(file_path: Path) -> dict[str, Any]:
    """
    Validate that file_path is an allowed Material Shortage Register workbook
    and contains all required columns.
    Returns column mapping if valid, or raises ValueError.
    """
    if not file_path.is_file():
        raise FileNotFoundError(f"Material Shortage Register file not found: {file_path}")

    if file_path.suffix.lower() not in (".xlsx", ".xls", ".xlsm"):
        raise ValueError(
            f"File '{file_path.name}' is not a valid Excel file. "
            f"Expected .xlsx, got '{file_path.suffix}'."
        )

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"Could not open '{file_path.name}' as Excel workbook: {exc}") from exc

    # Prefer sheet named 'Material Shortage Register' or active
    sheet = None
    for s_name in wb.sheetnames:
        if "shortage" in s_name.lower():
            sheet = wb[s_name]
            break
    if sheet is None:
        sheet = wb.active

    # Read header row
    header_row = next(sheet.iter_rows(values_only=True), None)
    wb.close()

    if not header_row:
        raise ValueError(f"File '{file_path.name}' is empty (no header row found).")

    # Map available column names (lowercase) to their 0-based column index
    col_indices: dict[str, int] = {}
    for idx, col in enumerate(header_row):
        if col is not None:
            col_indices[str(col).strip().casefold()] = idx

    # Check required fields
    resolved: dict[str, int] = {}
    missing_fields: list[str] = []

    for field_key, candidates in REQUIRED_REGISTER_FIELDS.items():
        found_idx = None
        for candidate in candidates:
            if candidate in col_indices:
                found_idx = col_indices[candidate]
                break
        if found_idx is not None:
            resolved[field_key] = found_idx
        else:
            missing_fields.append(candidates[0].title())

    if missing_fields:
        available_cols = [str(c) for c in header_row if c is not None]
        raise ValueError(
            f"File '{file_path.name}' is missing required column(s): {', '.join(missing_fields)}. "
            f"Found headers: {available_cols}"
        )

    # Optional date columns if present in register
    for date_key in ("vendor_date", "proposed_date", "final_date"):
        display_name = date_key.replace("_", " ")
        if display_name in col_indices:
            resolved[date_key] = col_indices[display_name]

    return resolved


def read_register_to_vendor_track_rows(
    file_path: Path,
    project_id: str,
    existing_dates: dict[tuple[str, str, str], tuple[str, str, str]] | None = None,
) -> list[tuple[Any, ...]]:
    """
    Read validated Material Shortage Register workbook and produce vendor_track rows.
    Preserves existing date values for matching (vendor, wbs, material) keys.
    """
    col_map = validate_material_shortage_register(file_path)

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = None
    for s_name in wb.sheetnames:
        if "shortage" in s_name.lower():
            sheet = wb[s_name]
            break
    if sheet is None:
        sheet = wb.active

    rows_iter = sheet.iter_rows(values_only=True)
    next(rows_iter, None)  # Skip header

    if existing_dates is None:
        existing_dates = {}

    v_idx = col_map["vendor_detail"]
    w_idx = col_map["wbs_element"]
    m_idx = col_map["material_code"]
    d_idx = col_map["description"]
    q_idx = col_map["quantity"]
    vd_idx = col_map.get("vendor_date")
    pd_idx = col_map.get("proposed_date")
    fd_idx = col_map.get("final_date")

    output: list[tuple[Any, ...]] = []
    for r in rows_iter:
        if not r or all(c is None for c in r):
            continue

        raw_vendor = clean_text(r[v_idx]) if v_idx < len(r) else ""
        wbs = clean_text(r[w_idx]) if w_idx < len(r) else ""
        material = clean_text(r[m_idx]) if m_idx < len(r) else ""
        desc = clean_text(r[d_idx]) if d_idx < len(r) else ""

        qty_val = r[q_idx] if q_idx < len(r) else 0.0
        try:
            qty = float(qty_val) if qty_val is not None else 0.0
        except (ValueError, TypeError):
            qty = 0.0

        # Vendor Detail is one ME2M vendor per shortage row. Do not split on
        # slash/pipe: those characters can be part of a legitimate vendor name.
        vendors = [raw_vendor] if raw_vendor else []

        for vendor in vendors:
            # Check if this item already had dates preserved in vendor_track
            lookup_key = (vendor, wbs, material)
            if lookup_key in existing_dates:
                v_date, p_date, f_date = existing_dates[lookup_key]
            else:
                v_date = clean_text(r[vd_idx]) if (vd_idx is not None and vd_idx < len(r)) else ""
                p_date = clean_text(r[pd_idx]) if (pd_idx is not None and pd_idx < len(r)) else ""
                f_date = clean_text(r[fd_idx]) if (fd_idx is not None and fd_idx < len(r)) else ""

            output.append((
                project_id,
                vendor,
                wbs,
                material,
                desc,
                qty,
                v_date,
                p_date,
                f_date,
                "",  # Status
            ))

    wb.close()
    return output


def refresh_vendor_track_from_registers(
    project_ids: list[str],
    registers_dir: Path = DEFAULT_SHORTAGE_REGISTERS_DIR,
    db_path: Path = SERVICES_DB_PATH,
) -> dict[str, Any]:
    """
    Process selected projects from their Material Shortage Register files in registers_dir.
    Validates all files first. Preserves existing date data for unchanged records.
    Idempotent: updates only selected projects without modifying other projects.
    """
    if not project_ids:
        raise ValueError("At least one project must be selected.")

    # Validate all files and resolve paths before touching database
    project_files: list[tuple[str, Path]] = []
    for pid in project_ids:
        clean_id = pid.strip()
        if not clean_id or not re.match(r"^[A-Za-z0-9][A-Za-z0-9\-\.]{0,63}$", clean_id):
            raise ValueError(f"Invalid project ID {pid!r}. Use only alphanumerics, hyphens, and dots.")

        # Strategy 1: Nested path — <registers_dir>/<project_id>/<project_id>_Net_shortage_Table.xlsx
        nested_file = registers_dir / clean_id / f"{clean_id}_Net_shortage_Table.xlsx"
        # Strategy 2: Flat path — <registers_dir>/<project_id>.xlsx
        flat_file = registers_dir / f"{clean_id}.xlsx"

        if nested_file.is_file():
            reg_file = nested_file
        elif flat_file.is_file():
            reg_file = flat_file
        else:
            # Try to find any xlsx in the nested folder
            nested_dir = registers_dir / clean_id
            reg_file = None
            if nested_dir.is_dir():
                for candidate in sorted(nested_dir.iterdir()):
                    if (
                        candidate.is_file()
                        and candidate.suffix.lower() in (".xlsx", ".xls", ".xlsm")
                        and not candidate.name.startswith("~$")
                    ):
                        reg_file = candidate
                        break
            if reg_file is None:
                raise FileNotFoundError(
                    f"Material Shortage Register file not found for project '{clean_id}': "
                    f"Expected '{nested_file.name}' in {registers_dir.name}/{clean_id}/ "
                    f"or '{flat_file.name}' in {registers_dir.name}/"
                )

        # Prevent path traversal
        if not str(reg_file.resolve()).startswith(str(registers_dir.resolve())):
            raise ValueError(f"Invalid register file path for project '{clean_id}'.")

        # Validate columns
        validate_material_shortage_register(reg_file)
        project_files.append((clean_id, reg_file))

    # Read existing dates from vendor_track for each project to preserve user dates
    all_columns = ["project_id"] + OUTPUT_COLUMNS
    columns_sql = ", ".join(quote_identifier(c) for c in all_columns)
    placeholders = ", ".join("?" for _ in all_columns)

    results: list[dict[str, Any]] = []

    with sqlite3.connect(db_path) as conn:
        conn.execute("PRAGMA busy_timeout = 5000")
        with conn:
            _ensure_vendor_track_schema(conn)

            for clean_id, reg_file in project_files:
                # Collect existing dates
                existing_dates: dict[tuple[str, str, str], tuple[str, str, str]] = {}
                cur = conn.execute(
                    '''SELECT "Vendor Detail", "WBS Element", "Material Code",
                              "Vendor Date", "Proposed Date", "Final Date"
                       FROM "vendor_track" WHERE "project_id" = ?''',
                    (clean_id,),
                )
                for row in cur.fetchall():
                    key = (clean_text(row[0]), clean_text(row[1]), clean_text(row[2]))
                    existing_dates[key] = (
                        clean_text(row[3]),
                        clean_text(row[4]),
                        clean_text(row[5]),
                    )

                # Build fresh rows from the register Excel file
                rows = read_register_to_vendor_track_rows(reg_file, clean_id, existing_dates)

                # Delete old rows for ONLY this project
                conn.execute(
                    f'DELETE FROM {quote_identifier(OUTPUT_TABLE)} WHERE "project_id" = ?',
                    (clean_id,),
                )

                # Insert fresh rows
                if rows:
                    conn.executemany(
                        f'INSERT INTO {quote_identifier(OUTPUT_TABLE)} ({columns_sql}) VALUES ({placeholders})',
                        rows,
                    )

                results.append({
                    "project_id": clean_id,
                    "rows": len(rows),
                    "filename": reg_file.name,
                })

    total_rows = sum(r["rows"] for r in results)
    return {
        "message": (
            f"Vendor Track processing complete for {len(results)} project(s). "
            f"{total_rows} vendor track records refreshed."
        ),
        "projects": results,
        "total_rows": total_rows,
    }


# ---------------------------------------------------------------------------
# Fallback / Database-direct paths
# ---------------------------------------------------------------------------
def build_rows_from_net_shortage(
    db_path: Path,
    project_id: str,
) -> list[tuple[Any, ...]]:
    """Read net_shortage from db_path and produce vendor_track rows tagged with project_id."""
    if not db_path.is_file():
        raise FileNotFoundError(f"Services database not found: {db_path}")
    required = ["Vendor Detail", "WBS Element", "Material Code", "Material Description", "Net Shortage"]
    with sqlite3.connect(db_path) as conn:
        conn.row_factory = sqlite3.Row
        table = resolve_table(conn, SOURCE_TABLE)
        columns = resolve_columns(conn, table, required)
        selection = ", ".join(
            f"{quote_identifier(columns[c])} AS {quote_identifier(c)}" for c in required
        )
        source_rows = conn.execute(f"SELECT {selection} FROM {quote_identifier(table)}").fetchall()

    output = []
    for row in source_rows:
        vendors = sorted({v.strip() for v in clean_text(row["Vendor Detail"]).split(" | ") if v.strip()})
        for vendor in vendors:
            output.append((
                project_id,
                vendor,
                clean_text(row["WBS Element"]),
                clean_text(row["Material Code"]),
                clean_text(row["Material Description"]),
                row["Net Shortage"] if row["Net Shortage"] is not None else 0,
                "", "", "", "",
            ))
    return output


def refresh_vendor_track_for_project(
    project_id: str,
    db_path: Path = SERVICES_DB_PATH,
) -> int:
    """Idempotent: delete existing rows for project_id, then insert fresh rows from net_shortage."""
    rows = build_rows_from_net_shortage(db_path, project_id)
    all_columns = ["project_id"] + OUTPUT_COLUMNS
    columns_sql = ", ".join(quote_identifier(c) for c in all_columns)
    placeholders = ", ".join("?" for _ in all_columns)

    with sqlite3.connect(db_path) as conn:
        conn.execute("PRAGMA busy_timeout = 5000")
        with conn:
            _ensure_vendor_track_schema(conn)
            conn.execute(
                f'DELETE FROM {quote_identifier(OUTPUT_TABLE)} WHERE "project_id" = ?',
                (project_id,),
            )
            if rows:
                conn.executemany(
                    f'INSERT INTO {quote_identifier(OUTPUT_TABLE)} ({columns_sql}) VALUES ({placeholders})',
                    rows,
                )
    print(f"Vendor tracking refreshed for project '{project_id}': {len(rows)} rows.")
    return len(rows)


def build_rows(db_path: Path = SERVICES_DB_PATH) -> list[tuple[Any, ...]]:
    '''Expand each Vendor Detail into a vendor-specific row (legacy, no project_id).'''
    if not db_path.is_file():
        raise FileNotFoundError(f"Services database not found: {db_path}")
    required = ["Vendor Detail", "WBS Element", "Material Code", "Material Description", "Net Shortage"]
    with sqlite3.connect(db_path) as conn:
        conn.row_factory = sqlite3.Row
        table = resolve_table(conn, SOURCE_TABLE)
        columns = resolve_columns(conn, table, required)
        selection = ", ".join(
            f"{quote_identifier(columns[c])} AS {quote_identifier(c)}" for c in required
        )
        source_rows = conn.execute(f"SELECT {selection} FROM {quote_identifier(table)}").fetchall()

    output = []
    for row in source_rows:
        vendors = sorted({v.strip() for v in clean_text(row["Vendor Detail"]).split(" | ") if v.strip()})
        for vendor in vendors:
            output.append((
                vendor,
                clean_text(row["WBS Element"]),
                clean_text(row["Material Code"]),
                clean_text(row["Material Description"]),
                row["Net Shortage"] if row["Net Shortage"] is not None else 0,
                "", "", "", "",
            ))
    return output


def refresh_vendor_track(db_path: Path = SERVICES_DB_PATH) -> int:
    '''Refresh vendor_track and preserve all other tables in services.db.'''
    rows = build_rows(db_path)
    definitions = (
        '"Vendor Detail" TEXT, "WBS Element" TEXT, "Material Code" TEXT, '
        '"Description" TEXT, "Quantity" REAL, "Vendor Date" TEXT, '
        '"Proposed Date" TEXT, "Final Date" TEXT, "Status" TEXT'
    )
    columns = ", ".join(quote_identifier(c) for c in OUTPUT_COLUMNS)
    placeholders = ", ".join("?" for _ in OUTPUT_COLUMNS)
    temp = "_vendor_track_new"
    with sqlite3.connect(db_path) as conn:
        conn.execute("PRAGMA busy_timeout = 5000")
        with conn:
            conn.execute(f"DROP TABLE IF EXISTS {quote_identifier(temp)}")
            conn.execute(f"CREATE TABLE {quote_identifier(temp)} ({definitions})")
            if rows:
                conn.executemany(
                    f"INSERT INTO {quote_identifier(temp)} ({columns}) VALUES ({placeholders})", rows
                )
            tables = {str(r[0]).casefold(): str(r[0]) for r in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            )}
            old = tables.get(OUTPUT_TABLE.casefold())
            if old:
                conn.execute(f"DROP TABLE {quote_identifier(old)}")
            conn.execute(f"ALTER TABLE {quote_identifier(temp)} RENAME TO {quote_identifier(OUTPUT_TABLE)}")
    print(f"Vendor tracking table refreshed: {OUTPUT_TABLE} ({len(rows)} rows)")
    return len(rows)


def main() -> int:
    try:
        print(f"Services database: {SERVICES_DB_PATH}")
        refresh_vendor_track(SERVICES_DB_PATH)
        return 0
    except Exception as exc:
        print(f"Vendor tracking service failed: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
