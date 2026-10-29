"""Django views over the existing SQLite and Python business services."""
from __future__ import annotations

import csv
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from urllib.parse import urlencode

from django.conf import settings
from django.core.paginator import Paginator
from django.http import FileResponse, Http404, HttpResponse, HttpResponseBadRequest
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils.text import get_valid_filename
from django.views.decorators.http import require_GET, require_POST

from shortage.data import API_COLUMNS, API_TO_DB, DB_COLUMNS, REPORTS

SAFE_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9-]{0,63}$")
REPORT_ORDER = ("zpp7", "mb25", "me2j", "pr", "me2m")
GLOBAL_ORDER = ("mb52", "scopelist")
CLASSIFIERS = {
    "zpp7": lambda p: "zpp7" in p.name.casefold(),
    "mb25": lambda p: "mb25" in p.name.casefold(),
    "me2j": lambda p: "me2j" in p.stem.casefold(),
    "pr": lambda p: p.stem.casefold() == "pr" or p.stem.casefold().startswith(("pr_", "pr ", "purchase requisition")),
    "me2m": lambda p: "me2m" in p.name.casefold(),
    "mb52": lambda p: "mb52" in p.name.casefold(),
    "scopelist": lambda p: "material_scope_list" in p.name.casefold(),
}


def _safe_id(value: str) -> str:
    if not SAFE_ID.fullmatch(value):
        raise Http404("Invalid project ID")
    return value


def _available_register_projects() -> list[str]:
    root = settings.SHORTAGE_REGISTERS_DIR.resolve()
    if not root.is_dir():
        return []
    projects = set()
    for folder in root.iterdir():
        if folder.is_dir() and SAFE_ID.fullmatch(folder.name):
            register = folder / f"{folder.name}_Net_shortage_Table.xlsx"
            if register.is_file():
                try:
                    register.resolve().relative_to(root)
                except ValueError:
                    continue
                projects.add(folder.name)
    for register in root.glob("*.xlsx"):
        if register.is_file() and SAFE_ID.fullmatch(register.stem):
            try:
                register.resolve().relative_to(root)
            except ValueError:
                continue
            projects.add(register.stem)
    return sorted(projects)


def _project_path(project_id: str) -> Path:
    project_id = _safe_id(project_id)
    root = settings.PROJECTS_DIR.resolve()
    path = (root / project_id).resolve()
    try:
        path.relative_to(root)
    except ValueError as exc:
        raise Http404("Project path is outside the configured project folder") from exc
    return path


def _find_reports(folder: Path, keys: tuple[str, ...]) -> dict[str, Path | None]:
    found = {key: None for key in keys}
    if folder.is_dir():
        for path in sorted(folder.iterdir()):
            if path.is_file() and not path.name.startswith("~$"):
                for key in keys:
                    if found[key] is None and CLASSIFIERS[key](path):
                        found[key] = path
                        break
    return found


def _conn(readonly=True):
    import sqlite3
    if not settings.SERVICES_DB.is_file():
        return None
    if readonly:
        conn = sqlite3.connect(settings.SERVICES_DB.as_uri() + "?mode=ro", uri=True)
    else:
        conn = sqlite3.connect(settings.SERVICES_DB)
    conn.row_factory = sqlite3.Row
    return conn


def _output_table(conn):
    names = {r[0].casefold(): r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    table = names.get("net_shortage")
    if not table:
        return None
    return '"' + table.replace('"', '""') + '"'


def _shortage_query(request, project_id=None):
    conn = _conn()
    if conn is None:
        return [], 0, {key: [] for key in ("material_ref_code", "wbs_element", "material_code", "scope")}, {}
    try:
        table = _output_table(conn)
        if table is None:
            return [], 0, {key: [] for key in ("material_ref_code", "wbs_element", "material_code", "scope")}, {}
        project_id = request.GET.get("project_id", "").strip() if project_id is None else project_id
        clauses, values = [], []
        if project_id and SAFE_ID.fullmatch(project_id):
            clauses.append('"project_id" = ?')
            values.append(project_id)
        filters = {}
        for field in ("material_ref_code", "wbs_element", "material_code", "scope"):
            col = f'"{API_TO_DB[field]}"'
            project_clause = '"project_id" = ? AND ' if project_id and SAFE_ID.fullmatch(project_id) else ""
            project_values = [project_id] if project_clause else []
            filters[field] = [r[0] for r in conn.execute(
                f"SELECT DISTINCT {col} FROM {table} WHERE {project_clause}"
                f"TRIM(COALESCE({col},''))<>'' ORDER BY {col}", project_values
            )]
        for field in ("material_ref_code", "wbs_element", "material_code", "scope"):
            value = request.GET.get(field, "").strip()
            if not value or value.casefold() == "all":
                continue
            col = f'"{API_TO_DB[field]}"'
            exact_values = {str(option).casefold() for option in filters[field]}
            if value.casefold() in exact_values:
                clauses.append(f"{col} = ? COLLATE NOCASE")
            else:
                clauses.append(f"instr(lower(COALESCE(CAST({col} AS TEXT), '')), lower(?)) > 0")
            values.append(value)
        where = " WHERE " + " AND ".join(clauses) if clauses else ""
        total = conn.execute(f"SELECT COUNT(*) FROM {table}{where}", values).fetchone()[0]
        order = '"Material Ref Code", "WBS Element", "Material Code"'
        rows = [dict(r) for r in conn.execute(
            f"SELECT \"project_id\", {', '.join(chr(34)+c+chr(34) for c in DB_COLUMNS)} "
            f"FROM {table}{where} ORDER BY {order}", values
        ).fetchall()]
        return rows, total, filters, {}
    finally:
        conn.close()


@require_GET
def shortages(request):
    projects = _available_register_projects()
    selected_project = request.GET.get("project_id", "").strip()
    if selected_project not in projects:
        selected_project = ""
    rows, total, filters, _ = _shortage_query(request, selected_project)
    page = Paginator(rows, 25).get_page(request.GET.get("page", 1))
    return render(request, "shortage/shortages.html", {"rows": page, "total": total, "filters": filters, "projects": projects, "selected_project": selected_project, "register_project_id": selected_project, "columns": DB_COLUMNS, "numeric_columns": {"ZPP7 Quantity", "Project Stock Quantity", "PO Quantity", "Open PO Quantity", "PR Quantity", "Net Shortage", "Free Stock"}})


@require_GET
def input_page(request):
    projects = sorted(p.name for p in settings.PROJECTS_DIR.iterdir() if p.is_dir() and SAFE_ID.fullmatch(p.name)) if settings.PROJECTS_DIR.exists() else []
    selected = request.GET.get("project", "")
    files = status = None
    if selected:
        project_dir = _project_path(selected)
        if not project_dir.is_dir():
            raise Http404("Project folder was not found")
        project_files = _find_reports(project_dir, REPORT_ORDER)
        global_files = _find_reports(settings.GLOBAL_DIR, GLOBAL_ORDER)
        files = [(key, REPORTS[key]["label"], project_files[key]) for key in REPORT_ORDER] + [(key, REPORTS[key]["label"], global_files[key]) for key in GLOBAL_ORDER]
        status = all(project_files[k] for k in REPORT_ORDER) and all(global_files[k] for k in GLOBAL_ORDER)
    register_path = None
    if selected:
        nested = settings.SHORTAGE_REGISTERS_DIR / selected / f"{selected}_Net_shortage_Table.xlsx"
        flat = settings.SHORTAGE_REGISTERS_DIR / f"{selected}.xlsx"
        register_path = nested if nested.is_file() else flat if flat.is_file() else None
    return render(request, "shortage/input.html", {"projects": projects, "selected": selected, "files": files, "ready": status, "register_exists": register_path is not None})


@require_POST
def upload_project_report(request, project_id):
    from django.contrib import messages
    project_id = _safe_id(project_id)
    kind = request.POST.get("report", "")
    upload = request.FILES.get("file")
    if kind not in REPORTS or upload is None:
        messages.error(request, "Choose a report type and a file to upload.")
        return redirect(f"/input/?project={project_id}")
    suffix = Path(upload.name).suffix.lower()
    if suffix not in REPORTS[kind]["extensions"]:
        messages.error(request, f"Unsupported extension for {REPORTS[kind]['label']}.")
        return redirect(f"/input/?project={project_id}")
    if upload.size == 0:
        messages.error(request, "The uploaded report is empty.")
        return redirect(f"/input/?project={project_id}")
    destination_dir = settings.GLOBAL_DIR.resolve() if kind in GLOBAL_ORDER else _project_path(project_id)
    destination_dir.mkdir(parents=True, exist_ok=True)
    safe_name = get_valid_filename(Path(upload.name).name)
    if not safe_name or safe_name in {".", ".."}:
        messages.error(request, "The uploaded filename is invalid.")
        return redirect(f"/input/?project={project_id}")
    destination = destination_dir / safe_name
    if destination.exists():
        messages.error(request, f"{safe_name} already exists. Rename the new file before uploading to preserve the existing report.")
        return redirect(f"/input/?project={project_id}")
    with destination.open("xb") as output:
        for chunk in upload.chunks():
            output.write(chunk)
    messages.success(request, f"Uploaded {safe_name}.")
    return redirect(f"/input/?project={project_id}")


@require_POST
def process_project(request, project_id):
    project_id = _safe_id(project_id)
    project_dir = _project_path(project_id)
    proj = _find_reports(project_dir, REPORT_ORDER)
    shared = _find_reports(settings.GLOBAL_DIR, GLOBAL_ORDER)
    missing = [k for k, v in {**proj, **shared}.items() if not v]
    if missing:
        from django.contrib import messages
        messages.error(request, "Missing required reports: " + ", ".join(missing))
        return redirect("input")
    settings.RAW_DIR.mkdir(parents=True, exist_ok=True)
    staged = {**proj, **shared}
    known_stems = {item["stem"].casefold() for item in REPORTS.values()}
    with tempfile.TemporaryDirectory(prefix="abb-raw-backup-") as backup_dir:
        backups = {}
        for existing in settings.RAW_DIR.iterdir():
            if existing.is_file() and (existing.stem.casefold() in known_stems or "me2m" in existing.stem.casefold()):
                saved = Path(backup_dir) / existing.name
                shutil.copy2(existing, saved)
                backups[existing.name] = saved
        try:
            for existing in settings.RAW_DIR.iterdir():
                if existing.is_file() and (existing.stem.casefold() in known_stems or "me2m" in existing.stem.casefold()):
                    existing.unlink()
            for key, source in staged.items():
                shutil.copy2(source, settings.RAW_DIR / f"{REPORTS[key]['stem']}{source.suffix}")
            from services.lock import processing_lock
            with processing_lock(project_id):
                child_env = os.environ.copy()
                child_env.update({
                    "ABB_DATA_ROOT": str(settings.DATA_ROOT),
                    "ABB_RAW_DIR": str(settings.RAW_DIR),
                    "ABB_GLOBAL_DB": str(settings.GLOBAL_DB),
                    "ABB_ZPP7_DB": str(settings.ZPP7_DB),
                    "ABB_SERVICES_DB": str(settings.SERVICES_DB),
                    "SHORTAGE_REGISTERS_DIR": str(settings.SHORTAGE_REGISTERS_DIR),
                    "VENDOR_REGISTERS_DIR": str(settings.VENDOR_REGISTERS_DIR),
                })
                result = subprocess.run([sys.executable, "main.py", "--report", "all", "--project-id", project_id], cwd=settings.BASE_DIR, capture_output=True, text=True, timeout=600, env=child_env)
        except Exception as exc:
            from django.contrib import messages
            messages.error(request, str(exc))
            return redirect("input")
        finally:
            for candidate in settings.RAW_DIR.iterdir():
                if candidate.is_file() and (candidate.stem.casefold() in known_stems or "me2m" in candidate.stem.casefold()):
                    candidate.unlink()
            for name, saved in backups.items():
                shutil.copy2(saved, settings.RAW_DIR / name)
    from django.contrib import messages
    if result.returncode:
        messages.error(request, (result.stdout + "\n" + result.stderr)[-6000:])
    else:
        from services.shortage_export import export_material_shortage_register
        export_material_shortage_register(project_id, db_path=settings.SERVICES_DB, output_dir=settings.SHORTAGE_REGISTERS_DIR)
        messages.success(request, f"Project {project_id} processed and its shortage register was updated.")
    return redirect("shortages")


@require_POST
def export_register(request, project_id):
    from django.contrib import messages
    try:
        from services.shortage_export import export_material_shortage_register
        info = export_material_shortage_register(_safe_id(project_id), db_path=settings.SERVICES_DB, output_dir=settings.SHORTAGE_REGISTERS_DIR)
        messages.success(request, f"Exported {info['filename']}.")
    except Exception as exc:
        messages.error(request, str(exc))
    return redirect("input")


@require_GET
def download_register(request, project_id):
    project_id = _safe_id(project_id)
    root = settings.SHORTAGE_REGISTERS_DIR.resolve()
    nested = root / project_id / f"{project_id}_Net_shortage_Table.xlsx"
    flat = root / f"{project_id}.xlsx"
    target = nested if nested.is_file() else flat
    target = target.resolve()
    try:
        target.relative_to(root)
    except ValueError as exc:
        raise Http404 from exc
    if not target.is_file():
        raise Http404("Register file was not found")
    return FileResponse(target.open("rb"), as_attachment=True, filename=target.name)


@require_POST
def open_project_register(request, project_id):
    from django.contrib import messages

    if request.META.get("REMOTE_ADDR") not in {"127.0.0.1", "::1"}:
        messages.error(request, "Opening a saved register is available from this computer only.")
        return redirect("shortages")

    project_id = _safe_id(project_id)
    root = settings.SHORTAGE_REGISTERS_DIR.resolve()
    nested = root / project_id / f"{project_id}_Net_shortage_Table.xlsx"
    flat = root / f"{project_id}.xlsx"
    target = (nested if nested.is_file() else flat).resolve()
    try:
        target.relative_to(root)
    except ValueError as exc:
        raise Http404 from exc
    if not target.is_file():
        messages.error(request, f"No saved shortage register was found for project {project_id}.")
        return redirect("shortages")
    if not hasattr(os, "startfile"):
        messages.error(request, "Direct Excel opening is available when the local app is running on Windows.")
        return redirect("shortages")

    try:
        os.startfile(str(target))
    except OSError as exc:
        messages.error(request, f"Could not open {target.name}: {exc}")
    else:
        messages.success(request, f"Opened saved register {target.name}.")
    return redirect("shortages")


def _latest_vendor_date_request(vendor):
    if not vendor or vendor.strip().casefold() in {"all", "all vendors"}:
        return None
    from services.vendor_request_export import sanitize_vendor_folder_name

    root = settings.VENDOR_REGISTERS_DIR.resolve()
    directory = (root / sanitize_vendor_folder_name(vendor) / "Date Request").resolve()
    try:
        directory.relative_to(root)
    except ValueError:
        return None
    if not directory.is_dir():
        return None
    candidates = []
    for path in directory.glob("*.xlsx"):
        if not path.is_file():
            continue
        target = path.resolve()
        try:
            target.relative_to(root)
        except ValueError:
            continue
        candidates.append(target)
    return max(candidates, key=lambda path: path.stat().st_mtime_ns, default=None)


def _latest_second_vendor_date_request(vendor):
    """Find the newest Second export inside the selected vendor's request folder."""
    if not vendor:
        return None
    from services.vendor_request_export import sanitize_vendor_folder_name

    root = settings.VENDOR_REGISTERS_DIR.resolve()
    directory = (root / sanitize_vendor_folder_name(vendor) / "Date Request").resolve()
    try:
        directory.relative_to(root)
    except ValueError:
        return None
    if not directory.is_dir():
        return None
    candidates = []
    for path in directory.iterdir():
        if not path.is_file() or path.suffix.casefold() not in {".xlsx", ".xlsm"}:
            continue
        if not re.fullmatch(r"Second(?:_v\d+)?", path.stem, flags=re.IGNORECASE):
            continue
        target = path.resolve()
        try:
            target.relative_to(root)
        except ValueError:
            continue
        candidates.append(target)
    return max(candidates, key=lambda path: path.stat().st_mtime_ns, default=None)


def _vendor_date_request_for_folder(vendor):
    """Prefer a Second request; fall back to another saved request for this vendor."""
    return _latest_second_vendor_date_request(vendor) or _latest_vendor_date_request(vendor)


@require_POST
def open_vendor_date_request(request):
    from django.contrib import messages

    vendor = request.POST.get("vendor", "").strip()
    query = {
        key: request.POST.get(key, "")
        for key in ("project_id", "vendor", "wbs_element")
        if request.POST.get(key, "")
    }
    return_to_vendor_track = lambda: redirect(f"/vendor-track/?{urlencode(query)}" if query else "vendor_track")

    if request.META.get("REMOTE_ADDR") not in {"127.0.0.1", "::1"}:
        messages.error(request, "Opening a saved vendor date request is available from this computer only.")
        return return_to_vendor_track()
    target = _latest_vendor_date_request(vendor)
    if target is None:
        messages.error(request, f"No exported vendor date request was found for {vendor or 'the selected vendor'}.")
        return return_to_vendor_track()
    if not hasattr(os, "startfile"):
        messages.error(request, "Direct Excel opening is available when the local app is running on Windows.")
        return return_to_vendor_track()
    try:
        from services.vendor_request_export import enforce_vendor_date_fields_editable

        enforce_vendor_date_fields_editable(target)
        os.startfile(str(target))
    except Exception as exc:
        messages.error(request, f"Could not open {target.name}: {exc}")
    else:
        messages.success(request, f"Opened the latest vendor date request: {target.name}.")
    return return_to_vendor_track()


@require_POST
def open_second_vendor_date_request(request):
    from django.contrib import messages

    vendor_folder = request.POST.get("vendor_folder", "").strip()
    if request.META.get("REMOTE_ADDR") not in {"127.0.0.1", "::1"}:
        messages.error(request, "Opening vendor date requests is available from this computer only.")
        return _vendor_response_redirect(vendor_folder)
    if not hasattr(os, "startfile"):
        messages.error(request, "Direct Excel opening is available when the local app is running on Windows.")
        return _vendor_response_redirect(vendor_folder)
    target = _vendor_date_request_for_folder(vendor_folder)
    if target is None:
        messages.error(request, f"No vendor date request was found for {vendor_folder or 'the selected vendor'}.")
        return _vendor_response_redirect(vendor_folder)
    try:
        from services.vendor_request_export import enforce_vendor_date_fields_editable

        enforce_vendor_date_fields_editable(target)
        os.startfile(str(target))
    except Exception as exc:
        messages.error(request, f"Could not open {target.name}: {exc}")
    else:
        messages.success(request, f"Opened vendor date request: {target.name}.")
    return _vendor_response_redirect(vendor_folder)


@require_POST
def delete_register(request, project_id):
    project_id = _safe_id(project_id)
    target = (settings.SHORTAGE_REGISTERS_DIR / project_id / f"{project_id}_Net_shortage_Table.xlsx").resolve()
    try:
        target.relative_to(settings.SHORTAGE_REGISTERS_DIR.resolve())
    except ValueError:
        raise Http404
    if target.is_file():
        target.unlink()
    return redirect("input")


@require_POST
def save_remarks(request):
    project, ref, wbs, material = (request.POST.get(k, "") for k in ("project_id", "material_ref_code", "wbs_element", "material_code"))
    if not project or not material:
        return HttpResponseBadRequest("Project and material identifiers are required.")
    conn = _conn(readonly=False)
    if conn is None:
        return HttpResponseBadRequest("Shortage database is not ready.")
    try:
        conn.execute('UPDATE net_shortage SET "Remarks"=? WHERE "project_id"=? AND "Material Ref Code"=? AND "WBS Element"=? AND "Material Code"=?', (request.POST.get("remarks", "")[:4000], project, ref, wbs, material))
        conn.commit()
    finally:
        conn.close()
    return redirect(request.META.get("HTTP_REFERER", "/"))


@require_GET
def export_shortages(request):
    rows, _, _, _ = _shortage_query(request)
    fmt = request.GET.get("format", "xlsx").lower()
    if fmt == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(DB_COLUMNS)
        writer.writerows([[r.get(c) for c in DB_COLUMNS] for r in rows])
        response = HttpResponse("\ufeff" + output.getvalue(), content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="abb-shortage-records.csv"'
        return response
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "Material Shortage Register"
    ws.append(DB_COLUMNS)
    header_fill = PatternFill("solid", fgColor="B5121B")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows: ws.append([row.get(c) for c in DB_COLUMNS])
    ws.freeze_panes = "A2"
    for index, width in enumerate((22, 22, 22, 40, 12, 16, 12, 14, 12, 14, 12, 30, 35, 30, 16), 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    stream = io.BytesIO(); wb.save(stream)
    response = HttpResponse(stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="abb-shortage-records.xlsx"'
    return response


@require_GET
def vendor_track(request):
    projects = _available_register_projects()
    conn = _conn()
    rows, vendors, wbs_options = [], [], []
    selected_project = request.GET.get("project_id", "")
    selected_vendor = request.GET.get("vendor", "")
    selected_wbs = request.GET.get("wbs_element", "")
    if conn:
        try:
            table = {r[0].casefold() for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            if "vendor_track" in table:
                columns = {r[1].casefold() for r in conn.execute('PRAGMA table_info("vendor_track")')}
                if "project_id" not in columns:
                    columns = set()
                if "project_id" in columns:
                    option_conditions, option_values = [], []
                    if selected_project:
                        option_conditions.append('"project_id" = ?'); option_values.append(selected_project)
                    option_where = " WHERE " + " AND ".join(option_conditions) if option_conditions else " WHERE 1=1"
                    vendors = [r[0] for r in conn.execute(f'SELECT DISTINCT "Vendor Detail" FROM vendor_track{option_where} AND TRIM(COALESCE("Vendor Detail",""))<>"" ORDER BY "Vendor Detail"', option_values)]
                    wbs_options = [r[0] for r in conn.execute(f'SELECT DISTINCT "WBS Element" FROM vendor_track{option_where} AND TRIM(COALESCE("WBS Element",""))<>"" ORDER BY "WBS Element"', option_values)]
                conditions, values = [], []
                for col, val, options in (("project_id", selected_project, projects), ("Vendor Detail", selected_vendor, vendors), ("WBS Element", selected_wbs, wbs_options)):
                    if not val or val.casefold() == "all":
                        continue
                    if val.casefold() in {str(option).casefold() for option in options}:
                        conditions.append(f'"{col}" = ? COLLATE NOCASE'); values.append(val)
                    else:
                        conditions.append(f"instr(lower(COALESCE(CAST(\"{col}\" AS TEXT), '')), lower(?)) > 0"); values.append(val)
                where = " WHERE " + " AND ".join(conditions) if conditions else ""
                if "project_id" in columns:
                    rows = [dict(r) for r in conn.execute(f'SELECT "project_id", "Vendor Detail", "WBS Element", "Material Code", "Description", "Quantity", "Vendor Date", "Proposed Date", "Final Date" FROM vendor_track{where} ORDER BY "Vendor Detail", "WBS Element", "Material Code"', values)]
        finally: conn.close()
    page = Paginator(rows, 50).get_page(request.GET.get("page", 1))
    vendor_is_selected = bool(selected_vendor and selected_vendor.casefold() in {str(vendor).casefold() for vendor in vendors})
    has_vendor_date_request = bool(_latest_vendor_date_request(selected_vendor)) if vendor_is_selected else False
    return render(request, "shortage/vendor_track.html", {"projects": projects, "records": page, "vendors": vendors, "wbs_options": wbs_options, "selected_project": selected_project, "selected_vendor": selected_vendor, "selected_wbs": selected_wbs, "register_project_id": selected_project if selected_project in projects else "", "vendor_is_selected": vendor_is_selected, "has_vendor_date_request": has_vendor_date_request})


@require_POST
def process_vendor_track(request):
    from django.contrib import messages
    from services.vendortrack_service import refresh_vendor_track_from_registers
    ids = request.POST.getlist("project_ids")
    if not ids:
        messages.error(request, "Select at least one processed project register to rebuild Vendor Track.")
        return redirect("vendor_track")
    try:
        refresh_vendor_track_from_registers(ids, registers_dir=settings.SHORTAGE_REGISTERS_DIR, db_path=settings.SERVICES_DB)
        messages.success(request, "Vendor Track refreshed.")
    except Exception as exc: messages.error(request, str(exc))
    return redirect("vendor_track")


@require_POST
def export_vendor_request(request):
    from django.contrib import messages
    from services.vendor_request_export import export_vendor_date_request
    try:
        result = export_vendor_date_request(vendor=request.POST.get("vendor", ""), project_ids=request.POST.getlist("project_ids"), wbs_elements=request.POST.getlist("wbs_elements"), db_path=settings.SERVICES_DB, output_root=settings.VENDOR_REGISTERS_DIR)
        target = Path(result["output_path"]).resolve()
        target.relative_to(settings.VENDOR_REGISTERS_DIR.resolve())
        messages.success(request, f"Vendor date request saved to {target.parent}\\{target.name}.")
        query = {
            "project_id": request.POST.get("project_ids", ""),
            "vendor": request.POST.get("vendor", ""),
            "wbs_element": request.POST.get("wbs_elements", ""),
        }
        query = {key: value for key, value in query.items() if value}
        return redirect(f"/vendor-track/?{urlencode(query)}" if query else "vendor_track")
    except Exception as exc: messages.error(request, str(exc))
    return redirect("vendor_track")


def _vendor_response_redirect(vendor_folder=""):
    url = reverse("vendor_import")
    if vendor_folder:
        url = f"{url}?{urlencode({'vendor_folder': vendor_folder})}"
    return redirect(url)


@require_GET
def vendor_import(request):
    from services.vendor_response_service import list_vendor_folders, read_vendor_register, resolve_vendor_workbook

    folders = list_vendor_folders(settings.VENDOR_DATE_REGISTERS_DIR)
    selected_folder = request.GET.get("vendor_folder", "")
    rows = []
    workbook_exists = False
    error = ""
    if selected_folder:
        if selected_folder not in folders:
            error = "Select a vendor folder from the list."
        else:
            try:
                workbook_path = resolve_vendor_workbook(settings.VENDOR_DATE_REGISTERS_DIR, selected_folder)
                rows = read_vendor_register(workbook_path)
                workbook_exists = True
            except Exception as exc:
                error = str(exc)
    accepted_count = sum(1 for row in rows if row["vendor_date_valid"])
    has_vendor_date_request = bool(_vendor_date_request_for_folder(selected_folder)) if workbook_exists else False
    return render(request, "shortage/vendor_import.html", {
        "folders": folders,
        "selected_folder": selected_folder,
        "rows": rows,
        "workbook_exists": workbook_exists,
        "accepted_count": accepted_count,
        "has_vendor_date_request": has_vendor_date_request,
    })


@require_POST
def open_vendor_register(request):
    from django.contrib import messages
    from services.vendor_response_service import resolve_vendor_workbook

    vendor_folder = request.POST.get("vendor_folder", "")
    if request.META.get("REMOTE_ADDR") not in {"127.0.0.1", "::1"}:
        messages.error(request, "Opening vendor workbooks is available from this computer only.")
        return _vendor_response_redirect(vendor_folder)
    if not hasattr(os, "startfile"):
        messages.error(request, "Direct Excel opening is available when the local app is running on Windows.")
        return _vendor_response_redirect(vendor_folder)
    try:
        target = resolve_vendor_workbook(settings.VENDOR_DATE_REGISTERS_DIR, vendor_folder)
        os.startfile(str(target))
        messages.success(request, f"Opened {target.name}.")
    except Exception as exc:
        messages.error(request, str(exc))
    return _vendor_response_redirect(vendor_folder)


@require_POST
def save_vendor_proposed_date(request):
    from django.contrib import messages
    from services.vendor_response_service import resolve_vendor_workbook, save_proposed_date

    vendor_folder = request.POST.get("vendor_folder", "")
    try:
        row_number = int(request.POST.get("row_number", ""))
        target = resolve_vendor_workbook(settings.VENDOR_DATE_REGISTERS_DIR, vendor_folder)
        save_proposed_date(target, row_number, request.POST.get("proposed_date", ""))
        messages.success(request, "Proposed Date saved to the vendor register.")
    except Exception as exc:
        messages.error(request, str(exc))
    return _vendor_response_redirect(vendor_folder)


@require_POST
def apply_vendor_dates(request):
    from django.contrib import messages
    from services.vendor_response_service import apply_accepted_vendor_dates, resolve_vendor_workbook

    vendor_folder = request.POST.get("vendor_folder", "")
    try:
        workbook_path = resolve_vendor_workbook(settings.VENDOR_DATE_REGISTERS_DIR, vendor_folder)
        result = apply_accepted_vendor_dates(
            workbook_path,
            db_path=settings.SERVICES_DB,
            registers_dir=settings.SHORTAGE_REGISTERS_DIR,
        )
        if result["updated_register_rows"]:
            messages.success(
                request,
                f"Applied {result['selected_rows']} Vendor Date row(s) to "
                f"{result['updated_register_rows']} material shortage row(s). "
                f"Regenerated registers for {len(result['projects'])} project(s).",
            )
        if result["missing_rows"]:
            messages.warning(request, f"No matching shortage rows were found for vendor workbook row(s): {', '.join(map(str, result['missing_rows']))}.")
        if result["export_errors"]:
            messages.error(request, "Final Dates were saved in the database, but some project registers could not be regenerated: " + "; ".join(result["export_errors"]))
        if not result["updated_register_rows"] and not result["missing_rows"]:
            messages.warning(request, "No vendor dates were applied.")
    except Exception as exc:
        messages.error(request, str(exc))
    return _vendor_response_redirect(vendor_folder)


@require_GET
def health(request):
    return HttpResponse("ok", content_type="text/plain")
